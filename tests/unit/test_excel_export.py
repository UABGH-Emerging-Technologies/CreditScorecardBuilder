"""Test suite for Excel export functionality."""

from io import BytesIO
from unittest.mock import MagicMock, Mock

import numpy as np
import openpyxl
import pandas as pd
import pytest
from openpyxl.worksheet.datavalidation import DataValidation

from CreditScore.excel_export import (
    ScorecardExcelExporter,
    create_scorecard_excel,
)

# Check if optbinning is available
try:
    import optbinning

    OPTBINNING_AVAILABLE = True
except ImportError:
    OPTBINNING_AVAILABLE = False


class TestScorecardExcelExporter:
    """Test cases for ScorecardExcelExporter class."""

    @pytest.fixture
    def mock_scorecard_model(self):
        """Create a mock scorecard model for testing."""
        model = Mock()

        # Create sample scorecard table data
        scorecard_data = pd.DataFrame(
            {
                "Variable": [
                    "feature1",
                    "feature1",
                    "feature1",
                    "feature2",
                    "feature2",
                    "Intercept",
                ],
                "Bin": ["(-inf, -0.5]", "(-0.5, 0.5]", "(0.5, inf)", "['A']", "['B' 'C']", "-"],
                "Points": [10, 20, 30, 15, 25, 50],
            }
        )

        model.get_scorecard_table.return_value = scorecard_data

        if OPTBINNING_AVAILABLE:
            # Create a proper mock scorecard that will pass isinstance check
            from optbinning import BinningProcess, Scorecard
            from sklearn.linear_model import LogisticRegression

            # Create sample data for fitting
            np.random.seed(42)
            X_fit = pd.DataFrame(
                {
                    "feature1": np.random.randn(100),
                    "feature2": np.random.choice(["A", "B", "C"], 100),
                }
            )
            y_fit = np.random.binomial(1, 0.3, 100)

            # Create and fit a minimal real scorecard for monitoring
            binning_process = BinningProcess(variable_names=["feature1", "feature2"])
            binning_process.fit(X_fit, y_fit)

            mock_scorecard = Scorecard(
                binning_process=binning_process,
                estimator=LogisticRegression(),
                scaling_method="min_max",
                scaling_method_params={"min": 0, "max": 100},
            )
            mock_scorecard.fit(X_fit, y_fit)

            model.scorecard = mock_scorecard
        else:
            # Fallback for when optbinning is not available
            model.scorecard = Mock()

        return model

    @pytest.fixture
    def sample_data(self):
        """Create sample X and y data for monitoring."""
        np.random.seed(42)
        X = pd.DataFrame(
            {"feature1": np.random.randn(100), "feature2": np.random.choice(["A", "B", "C"], 100)}
        )
        y = np.random.binomial(1, 0.3, 100)
        return X, y

    @pytest.fixture
    def exporter(self):
        """Create an exporter instance."""
        return ScorecardExcelExporter()

    def test_parse_categorical_bin(self, exporter):
        """Test parsing of categorical bin strings."""
        # Test single value
        assert exporter._parse_categorical_bin("['Male']") == ["Male"]

        # Test multiple values
        assert exporter._parse_categorical_bin("['A' 'B' 'C']") == ["A", "B", "C"]

        # Test special values
        assert exporter._parse_categorical_bin("Missing") == []
        assert exporter._parse_categorical_bin("Special") == []

        # Test non-array format
        assert exporter._parse_categorical_bin("SimpleValue") == ["SimpleValue"]

    def test_format_bin_display(self, exporter):
        """Test formatting of bin strings for display."""
        # Test numerical bins
        assert exporter._format_bin_display("(-inf, 30.0]") == "≤ 30.0"
        assert exporter._format_bin_display("[50.0, inf)") == "≥ 50.0"

        # Test categorical bins
        assert exporter._format_bin_display("['Male']") == "Male"
        assert exporter._format_bin_display("['A' 'B']") == "A / B"

        # Test special cases
        assert exporter._format_bin_display("Missing") == "Missing"
        assert exporter._format_bin_display("Special") == "Special"

        # Test regular ranges (should use clean mathematical notation)
        assert exporter._format_bin_display("[2.5, 5.27]") == "2.5 ≤ x ≤ 5.27"
        assert exporter._format_bin_display("(2.5, 5.27)") == "2.5 < x < 5.27"

    def test_export_scorecard_creates_all_sheets(self, mock_scorecard_model, sample_data):
        """Test that export creates all required sheets."""
        exporter = ScorecardExcelExporter()
        feature_names = ["feature1", "feature2"]
        X, y = sample_data

        # Export to BytesIO
        excel_bytes = exporter.export_scorecard(mock_scorecard_model, feature_names, X, y)

        # Load the workbook and check sheets
        wb = openpyxl.load_workbook(excel_bytes)
        assert "Scorecard Calculator" in wb.sheetnames
        assert "Lookups" in wb.sheetnames
        assert "PSI Lookup" in wb.sheetnames

        # Verify sheets are not hidden
        lookups_sheet = wb["Lookups"]
        psi_sheet = wb["PSI Lookup"]
        assert lookups_sheet.sheet_state != "hidden"
        assert psi_sheet.sheet_state != "hidden"

    def test_export_scorecard_creates_data_validations(self, mock_scorecard_model, sample_data):
        """Test that data validations (dropdowns) are created."""
        exporter = ScorecardExcelExporter()
        feature_names = ["feature1", "feature2"]
        X, y = sample_data

        excel_bytes = exporter.export_scorecard(mock_scorecard_model, feature_names, X, y)

        # Load and check for data validations
        wb = openpyxl.load_workbook(excel_bytes)
        main_sheet = wb["Scorecard Calculator"]

        # Should have data validations
        assert len(main_sheet.data_validations) > 0

        # Check that validations are of type 'list'
        for dv in main_sheet.data_validations.dataValidation:
            assert dv.type == "list"
            assert dv.allow_blank is True

    def test_export_scorecard_unlocked_input_cells(self, mock_scorecard_model, sample_data):
        """Test that input cells are unlocked for editing."""
        exporter = ScorecardExcelExporter()
        feature_names = ["feature1", "feature2"]
        X, y = sample_data

        excel_bytes = exporter.export_scorecard(mock_scorecard_model, feature_names, X, y)

        # Load and check cell protection
        wb = openpyxl.load_workbook(excel_bytes)
        main_sheet = wb["Scorecard Calculator"]

        # Check that sheet is protected
        assert main_sheet.protection.sheet is True

        # Check cells that have data validation (these should be unlocked)
        validation_cells = []
        for dv in main_sheet.data_validations.dataValidation:
            validation_cells.extend(dv.cells)

        # Verify that at least some cells have validation
        assert len(validation_cells) > 0, "No validation cells found"

        # Check that validation cells are unlocked
        for cell_range in validation_cells:
            for cell in cell_range:
                if hasattr(cell, "protection") and cell.protection:
                    assert (
                        cell.protection.locked is False
                    ), f"Cell {cell.coordinate} should be unlocked"

    def test_export_scorecard_creates_vlookup_formulas(self, mock_scorecard_model, sample_data):
        """Test that VLOOKUP formulas are created correctly."""
        exporter = ScorecardExcelExporter()
        feature_names = ["feature1", "feature2"]
        X, y = sample_data

        excel_bytes = exporter.export_scorecard(mock_scorecard_model, feature_names, X, y)

        # Load and check formulas
        wb = openpyxl.load_workbook(excel_bytes)
        main_sheet = wb["Scorecard Calculator"]

        # Find cells with VLOOKUP formulas
        vlookup_count = 0
        psi_lookup_refs = 0

        for row in main_sheet.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    if "VLOOKUP" in cell.value:
                        vlookup_count += 1
                        # Verify formula references correct sheets
                        if "Lookups!" in cell.value:
                            # Standard feature lookup
                            pass
                        elif "'PSI Lookup'!" in cell.value:
                            # PSI lookup for probability/risk
                            psi_lookup_refs += 1

        assert vlookup_count > 0, "No VLOOKUP formulas found in the worksheet"
        assert psi_lookup_refs > 0, "No PSI Lookup references found"

    def test_export_scorecard_with_output_file(self, mock_scorecard_model, sample_data, tmp_path):
        """Test saving to a file path."""
        exporter = ScorecardExcelExporter()
        feature_names = ["feature1", "feature2"]
        X, y = sample_data
        output_file = tmp_path / "test_output.xlsx"

        excel_bytes = exporter.export_scorecard(
            mock_scorecard_model, feature_names, X, y, output_file=str(output_file)
        )

        # Verify file was created
        assert output_file.exists()

        # Verify BytesIO is still at position 0
        assert excel_bytes.tell() == 0

    def test_create_scorecard_excel_function(self, mock_scorecard_model, sample_data):
        """Test the convenience function."""
        feature_names = ["feature1", "feature2"]
        X, y = sample_data

        result = create_scorecard_excel(mock_scorecard_model, feature_names, X, y)

        # Should return BytesIO
        assert isinstance(result, BytesIO)

        # Should be at position 0
        assert result.tell() == 0

        # Should be a valid Excel file
        wb = openpyxl.load_workbook(result)
        assert len(wb.sheetnames) >= 3  # Main, Lookups, PSI Lookup


class TestExcelExportIntegration:
    """Integration tests with actual optbinning models."""

    @pytest.mark.skipif(not pytest.importorskip("optbinning"), reason="optbinning not installed")
    def test_with_real_scorecard_model(self):
        """Test with a real OptbinningScorecardModel."""
        from optbinning import BinningProcess, Scorecard
        from sklearn.linear_model import LogisticRegression

        # Create sample data
        np.random.seed(42)
        n_samples = 100
        X = pd.DataFrame(
            {
                "feature1": np.random.randn(n_samples),
                "feature2": np.random.choice(["A", "B", "C"], n_samples),
            }
        )
        y = np.random.binomial(1, 0.3, n_samples)

        # Create binning process and scorecard
        binning_process = BinningProcess(variable_names=["feature1", "feature2"])
        binning_process.fit(X, y)

        scorecard = Scorecard(
            binning_process=binning_process,
            estimator=LogisticRegression(),
            scaling_method="min_max",
            scaling_method_params={"min": 0, "max": 100},
        )
        scorecard.fit(X, y)

        # Mock the get_scorecard_table method since it doesn't exist on base Scorecard
        scorecard.get_scorecard_table = lambda: scorecard.table()

        # Test export
        excel_bytes = create_scorecard_excel(scorecard, ["feature1", "feature2"], X, y)

        # Verify it's a valid Excel file with expected content
        wb = openpyxl.load_workbook(excel_bytes)
        assert "Scorecard Calculator" in wb.sheetnames
        assert "Lookups" in wb.sheetnames
